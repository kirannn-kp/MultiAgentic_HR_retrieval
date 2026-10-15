"""Import Name/Region/Description/URL from Excel into local approved-source configuration; no Azure calls."""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from mbpoc.regional import ApprovedSource, REGIONS, SourceCatalog  # noqa: E402


def read_catalog(path: Path) -> SourceCatalog:
    book = load_workbook(path, read_only=True, data_only=True)
    sources = []
    aliases = {"Germany": "DE", "European Union": "EU", "DE": "DE", "EU": "EU"}
    try:
        for sheet in book:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, ())
            columns = {str(value).strip(): index for index, value in enumerate(headers) if value is not None}
            if not {"Name", "Region", "Description", "URL"} <= columns.keys():
                raise ValueError(f"Sheet '{sheet.title}' must have Name, Region, Description and URL headers")
            for number, values in enumerate(rows, 2):
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                data = {key: values[index] for key, index in columns.items()}
                if data["Region"] not in aliases:
                    raise ValueError(f"{sheet.title} row {number}: unsupported Region '{data['Region']}' (use Germany or European Union)")
                sources.append(ApprovedSource(
                    name=data["Name"], region=aliases[data["Region"]],
                    description=data["Description"], url=data["URL"],
                ))
    finally:
        book.close()
    catalog = SourceCatalog(sources=sources)
    catalog.web_definitions()
    return catalog


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=ROOT / "config" / "approved_web_sources.yaml")
    args = parser.parse_args()
    catalog = read_catalog(args.workbook)
    args.output.write_text(yaml.safe_dump(catalog.model_dump(), allow_unicode=True, sort_keys=False), encoding="utf-8")
    for region, label in REGIONS.items():
        print(f"{label}: {len(catalog.group(region))} approved websites")
    print(f"Saved local catalogue to {args.output}. No Azure resources were changed.")


if __name__ == "__main__":
    main()
