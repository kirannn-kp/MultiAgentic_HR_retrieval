"""'SharePoint index' without a cross-tenant indexer app: read the documents (from a SharePoint site via Microsoft Graph,
or from data/sharepoint) and push them into an Azure AI Search index that Foundry IQ uses as the `sharepoint` knowledge source.

Needs only 'Search Service Contributor' + 'Search Index Data Contributor' on the Search service (Entra auth).
"""
from __future__ import annotations

import hashlib
import io
import os
import re
from pathlib import Path
from typing import Iterator
from urllib.parse import urlparse

import requests
from azure.identity import DefaultAzureCredential

from ..config import KS_SHAREPOINT, ROOT, Settings

INDEX_NAME = "mb-sharepoint-docs"
SEMANTIC_CONFIG = "default"
READ_SCOPES = ["https://graph.microsoft.com/Sites.Read.All", "https://graph.microsoft.com/Files.Read.All", "https://graph.microsoft.com/User.Read"]
SUPPORTED = {".md", ".txt", ".docx", ".pdf", ".html", ".csv"}
_cred = DefaultAzureCredential()


def _hdr() -> dict:
    return {"Authorization": f"Bearer {_cred.get_token('https://search.azure.com/.default').token}", "Content-Type": "application/json"}


def _url(s: Settings, path: str) -> str:
    return f"{s.search_endpoint.rstrip('/')}/{path}{'&' if '?' in path else '?'}api-version={s.search_api_version}"


def ensure_index(s: Settings) -> None:
    body = {
        "name": INDEX_NAME,
        "fields": [
            {"name": "id", "type": "Edm.String", "key": True, "filterable": True},
            {"name": "title", "type": "Edm.String", "searchable": True, "retrievable": True},
            {"name": "content", "type": "Edm.String", "searchable": True, "retrievable": True},
            {"name": "source_url", "type": "Edm.String", "retrievable": True, "filterable": True},
            {"name": "path", "type": "Edm.String", "retrievable": True, "filterable": True},
        ],
        "semantic": {"defaultConfiguration": SEMANTIC_CONFIG, "configurations": [{
            "name": SEMANTIC_CONFIG,
            "prioritizedFields": {"titleField": {"fieldName": "title"}, "prioritizedContentFields": [{"fieldName": "content"}]},
        }]},
    }
    resp = requests.put(_url(s, f"indexes/{INDEX_NAME}"), headers=_hdr(), json=body, timeout=60)
    if resp.status_code >= 400:
        raise RuntimeError(f"create index -> {resp.status_code}: {resp.text[:600]}")


def chunk(text: str, size: int = 900) -> list[str]:
    """Split on markdown headings / paragraphs into ~size-character chunks (oversized paragraphs are split on sentences)."""
    parts: list[str] = []
    for p in re.split(r"\n(?=#{1,3}\s)|\n\s*\n", text):
        p = p.strip()
        while len(p) > size:
            cut = max(p.rfind(". ", 0, size), p.rfind("\n", 0, size))
            cut = cut + 1 if cut > size // 2 else size
            parts.append(p[:cut].strip())
            p = p[cut:].strip()
        if p:
            parts.append(p)
    chunks, cur = [], ""
    for p in parts:
        if cur and len(cur) + len(p) > size:
            chunks.append(cur)
            cur = p
        else:
            cur = f"{cur}\n\n{p}" if cur else p
    if cur:
        chunks.append(cur)
    return chunks


def extract_text(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()
    if ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        tables = ["\n".join(" | ".join(c.text for c in row.cells) for row in t.rows) for t in d.tables]
        return "\n\n".join([p.text for p in d.paragraphs if p.text.strip()] + tables)
    if ext == ".pdf":
        from pypdf import PdfReader

        return "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    if ext == ".html":
        return re.sub(r"<[^>]+>", " ", data.decode("utf-8", "ignore"))
    return data.decode("utf-8", "ignore")


def local_documents(base_url: str, folder: Path | None = None) -> Iterator[dict]:
    folder = folder or ROOT / "data" / "sharepoint"
    for path in sorted(p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED):
        rel = path.relative_to(folder).as_posix()
        yield {"name": path.name, "path": rel, "url": f"{base_url}/Shared%20Documents/{rel}", "data": path.read_bytes()}


def graph_documents(s: Settings) -> tuple[str, Iterator[dict]]:
    from .graph import AzCliGraph, Graph, site_from_url

    auth = os.environ.get("GRAPH_AUTH", "auto").lower()  # az | device | auto
    g = None
    if auth in ("az", "auto"):
        try:
            g = AzCliGraph(s.tenant_id)
            site = site_from_url(g, s.sharepoint_site_url)
            print("Signed in through the Azure CLI account of the SharePoint tenant")
        except Exception as exc:
            if auth == "az":
                raise
            print(f"(Azure CLI route not available: {str(exc)[:160]}) -> device-code sign-in")
            g = None
    if g is None:
        g = Graph(s.tenant_id, READ_SCOPES)
        site = site_from_url(g, s.sharepoint_site_url)
    drives = g.request("GET", f"/sites/{site['id']}/drives").json()["value"]

    def walk(drive_id: str, item_id: str, prefix: str) -> Iterator[dict]:
        url = f"/drives/{drive_id}/items/{item_id}/children?$top=200"
        while url:
            page = g.request("GET", url).json()
            for it in page["value"]:
                if "folder" in it:
                    yield from walk(drive_id, it["id"], f"{prefix}{it['name']}/")
                elif Path(it["name"]).suffix.lower() in SUPPORTED:
                    link = it.get("@microsoft.graph.downloadUrl")
                    try:
                        if link:
                            resp = requests.get(link, timeout=300)
                            resp.raise_for_status()
                            data = resp.content
                        else:
                            data = g.request("GET", f"/drives/{drive_id}/items/{it['id']}/content").content
                    except Exception as exc:
                        print(f"  ! cannot download {prefix}{it['name']} ({str(exc)[:80]}) - the tenant blocks file content for this sign-in; "
                              "download the file in the browser and use: python scripts/01_sync_sharepoint.py --folder <dir>")
                        continue
                    yield {"name": it["name"], "path": f"{prefix}{it['name']}", "url": it.get("webUrl", ""), "data": data}
            url = page.get("@odata.nextLink")

    def all_docs() -> Iterator[dict]:
        for d in drives:
            yield from walk(d["id"], "root", "")

    return site["webUrl"], all_docs()


def sync(s: Settings, source: str = "auto", folder: str | None = None) -> int:
    """source: graph (SharePoint site via Microsoft Graph), folder (local export of the site), local (sample docs) or auto
    (graph if SHAREPOINT_SITE_URL is set, otherwise the sample docs)."""
    use_graph = source == "graph" or (source == "auto" and bool(s.sharepoint_site_url) and not folder)
    if use_graph:
        base_url, docs = graph_documents(s)
        print(f"Reading SharePoint site {base_url} via Microsoft Graph ...")
    elif folder:
        base_url = s.sharepoint_site_url or "https://sharepoint.example/sites/MB-Enterprise-POC"
        docs = local_documents(base_url, Path(folder))
        print(f"Using documents exported to {folder}")
    else:
        base_url = s.sharepoint_site_url or "https://sharepoint.example/sites/MB-Enterprise-POC"
        docs = local_documents(base_url)
        print("Using sample documents from data/sharepoint (no SharePoint tenant access)")
    ensure_index(s)
    batch: list[dict] = []
    for doc in docs:
        text = extract_text(doc["name"], doc["data"])
        first = re.search(r"^#\s+(.+)$", text, re.M)
        title = first.group(1).strip() if first else Path(doc["name"]).stem.replace("-", " ").replace("_", " ")
        pieces = chunk(text)
        print(f"  {doc['path']}: {len(pieces)} chunks")
        for i, c in enumerate(pieces):
            key = hashlib.sha1(f"{doc['path']}#{i}".encode()).hexdigest()
            batch.append({"@search.action": "mergeOrUpload", "id": key, "title": title, "content": c, "source_url": doc["url"], "path": doc["path"]})
    if not batch:
        raise SystemExit("No documents could be read. If the tenant blocks file download for this sign-in, download the files in the browser and run --folder <dir>.")
    for i in range(0, len(batch), 100):
        resp = requests.post(_url(s, f"indexes/{INDEX_NAME}/docs/index"), headers=_hdr(), json={"value": batch[i:i + 100]}, timeout=120)
        if resp.status_code >= 400:
            raise RuntimeError(f"upload -> {resp.status_code}: {resp.text[:600]}")
    print(f"Indexed {len(batch)} chunks into '{INDEX_NAME}'")
    removed = _prune(s, {d["id"] for d in batch})
    if removed:
        print(f"Removed {removed} stale chunks (documents no longer in the source)")
    return len(batch)


def _prune(s: Settings, keep: set[str]) -> int:
    """Make the index mirror the source: delete chunks that are not part of the current sync."""
    resp = requests.post(_url(s, f"indexes/{INDEX_NAME}/docs/search"), headers=_hdr(), json={"search": "*", "select": "id", "top": 1000}, timeout=60)
    resp.raise_for_status()
    stale = [d["id"] for d in resp.json().get("value", []) if d["id"] not in keep]
    for i in range(0, len(stale), 100):
        requests.post(_url(s, f"indexes/{INDEX_NAME}/docs/index"), headers=_hdr(),
                      json={"value": [{"@search.action": "delete", "id": k} for k in stale[i:i + 100]]}, timeout=60).raise_for_status()
    return len(stale)


def knowledge_source_body() -> dict:
    return {
        "name": KS_SHAREPOINT, "kind": "searchIndex",
        "description": "HR / compliance policies from the SharePoint library (synced into Azure AI Search)",
        "searchIndexParameters": {
            "searchIndexName": INDEX_NAME, "semanticConfigurationName": SEMANTIC_CONFIG,
            "sourceDataFields": [{"name": "title"}, {"name": "source_url"}, {"name": "path"}],
        },
    }
